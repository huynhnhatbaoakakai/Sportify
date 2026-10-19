from pathlib import Path
import re, json, zipfile, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

root=Path(__file__).parent
source=root/'SPORTIFY_DOC_DE_HIEU.md'
text=source.read_text(encoding='utf-8')
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8)
section.left_margin=Cm(1.9);section.right_margin=Cm(1.9)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
for name,size in [('Title',23),('Heading 1',16),('Heading 2',13)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('174E65');st.paragraph_format.keep_with_next=True
if 'Source Code' not in doc.styles:
 st=doc.styles.add_style('Source Code',1)
else:st=doc.styles['Source Code']
st.font.name='Consolas';st.font.size=Pt(9)
st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.0
st.paragraph_format.left_indent=Cm(.25)
header=section.header.paragraphs[0];header.text='SPORTIFY  |  ĐỌC HIỂU SOURCE CODE';header.style=doc.styles['Caption']
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Sportify • Tài liệu học, không phải chứng nhận runtime • Trang ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)

def inline(paragraph,value):
 for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',value):
  if not part:continue
  run=paragraph.add_run(part[2:-2] if part.startswith('**') and part.endswith('**') else part[1:-1] if part.startswith('`') and part.endswith('`') else part)
  if part.startswith('**') and part.endswith('**'):run.bold=True
  if part.startswith('`') and part.endswith('`'):run.font.name='Consolas';run.font.size=Pt(10)

def table(rows):
 t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,v in enumerate(rows[0]):inline(t.rows[0].cells[i].paragraphs[0],v)
 prop=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');prop.append(repeat)
 for vals in rows[1:]:
  cells=t.add_row().cells
  for j,v in enumerate(vals):inline(cells[j].paragraphs[0],v)
 for row in t.rows:
  for cell in row.cells:
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(4)
    for r in par.runs:r.font.size=Pt(10)
 doc.add_paragraph()

lines=text.splitlines();i=0;headings=[];code_blocks=0
while i<len(lines):
 line=lines[i]
 if line.startswith('```'):
  code=[];i+=1
  while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
  par=doc.add_paragraph('\n'.join(code),'Source Code')
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F8');par._p.get_or_add_pPr().append(sh)
  code_blocks+=1
 elif line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in vals):rows.append(vals)
   i+=1
  table(rows);continue
 elif line.startswith('# '):
  doc.add_heading(line[2:],0)
 elif line.startswith('### '):
  doc.add_heading(line[4:],2)
 elif line.startswith('## '):
  doc.add_heading(line[3:],1);headings.append(line[3:])
 elif line.strip():
  inline(doc.add_paragraph(),line)
 i+=1

doc.core_properties.title='Sportify — Đọc để hiểu và tự giải thích code'
doc.core_properties.subject='Luồng Servlet/JSP, JPA/JDBC, owner session và tự xác nhận đơn mới'
doc.core_properties.author='Tài liệu ôn tập'
out=root/'SPORTIFY_DOC_DE_HIEU.docx';doc.save(out)
# Verify real package/readback; this does not claim rendered Word pagination.
readback=Document(out)
full='\n'.join(par.text for par in readback.paragraphs)+'\n'+'\n'.join(c.text for t in readback.tables for r in t.rows for c in r.cells)
required=['Đơn A vẫn PENDING','autoApprove KHÔNG có nghĩa là PENDING','getResultList','setAutoApproval','createCustomerBooking','OwnerAccess','PreparedStatement','getAttribute','SportFieldDB']
missing=[s for s in required if s not in full]
with zipfile.ZipFile(out) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
assert len(headings)==16,len(headings)
assert not missing,missing
report={'artifact':str(out),'bytes':out.stat().st_size,'main_sections':len(headings),'tables':len(readback.tables),'code_blocks':code_blocks,'required_content_missing':missing,'docx_readback':'PASS','zip_xml_health':'PASS','visual_render':'Not performed; no rendered page-count claim.'}
(root/'SPORTIFY_DOC_DE_HIEU_VERIFY.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
