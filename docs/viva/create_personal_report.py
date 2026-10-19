from pathlib import Path
import json, zipfile, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

outdir=Path(__file__).parent
out=outdir/'BAO_CAO_CA_NHAN_HUYNH_NHAT_BAO_24110079.docx'
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
n=doc.styles['Normal'];n.font.name='Times New Roman';n.font.size=Pt(12)
n.paragraph_format.line_spacing=1.05;n.paragraph_format.space_after=Pt(5)
for name in ['Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('17365D');st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(4)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('BÁO CÁO THỰC HÀNH CÁ NHÂN');r.bold=True;r.font.size=Pt(15)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Bật/tắt chế độ tự xác nhận đơn mới cho sân');r.bold=True;r.font.size=Pt(13)
for label,value in [('Sinh viên','Huỳnh Nhật Bảo'),('MSSV','24110079'),('Môn học','Web Programming'),('Dự án','Sportify')]:
 p=doc.add_paragraph();p.add_run(label+': ').bold=True;p.add_run(value)

def h(t):doc.add_heading(t,level=1)
def p(t):doc.add_paragraph(t)
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for r in t.rows[0].cells[0].paragraphs[0].runs:r.bold=True
 for cell in t.rows[0].cells:
  for run in cell.paragraphs[0].runs:run.bold=True
 pr=t.rows[0]._tr.get_or_add_trPr();pr.append(OxmlElement('w:tblHeader'))
 for vals in rows:
  cells=t.add_row().cells
  for i,x in enumerate(vals):cells[i].text=x
 for row in t.rows:
  for cell in row.cells:
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(3)
    for run in par.runs:run.font.size=Pt(10.5)

h('1. Mục tiêu và phạm vi thực hiện')
p('Trong buổi thực hành tại lớp, em thực hiện chức năng nhỏ cho phép chủ sân bật hoặc tắt chế độ tự xác nhận đơn đặt sân mới. Chức năng giúp chủ sân lựa chọn duyệt thủ công hoặc cho hệ thống xác nhận đơn mới hợp lệ theo cấu hình của từng sân.')
p('Phạm vi báo cáo là thao tác thay đổi và lưu cấu hình của sân; không bao gồm toàn bộ chức năng Booking, kiểm tra trùng lịch và tạo đơn do thành viên khác phụ trách. Công nghệ liên quan: JSP, Jakarta Servlet, JPA/Hibernate và PostgreSQL.')

h('2. Thiết kế và vị trí mã nguồn')
p('Mỗi sân có thuộc tính boolean autoApprove trong model SportField, ánh xạ tới cột auto_approve của bảng sports_fields. Giá trị true là bật, false là tắt. Quy tắc hiện tại chỉ cho bật với sân bóng rổ và pickleball thông qua supportsAutoApproval().')
table(['Thành phần','Vị trí / vai trò'],[
('Giao diện','WEB-INF/views/owner/calendar.jsp: hiển thị chế độ và form bật/tắt.'),
('Servlet','controller/OwnerCalendarServlet.java, doPost(): nhận POST /owner/fields/approval.'),
('Truy cập dữ liệu','data/SportFieldDB.java, setAutoApproval(): kiểm tra và lưu cấu hình.'),
('Model','model/SportField.java: autoApprove, isAutoApprove(), setAutoApprove(), supportsAutoApproval().')])
p('Đường dẫn JSP tính từ src/main/webapp/; các đường dẫn Java tính từ src/main/java/com/sportify/.')

h('3. Luồng xử lý')
for t in [
'Chủ sân chọn sân tại trang lịch và bấm nút bật/tắt. Form gửi POST với fieldId, enabled, date và csrfToken. Giá trị enabled là trạng thái ngược với cấu hình hiện tại.',
'OwnerCalendarServlet lấy ownerId từ tài khoản trong session, kiểm tra sân thuộc chủ sân và chỉ chấp nhận enabled là chuỗi "true" hoặc "false". Yêu cầu owner được kiểm tra quyền và token CSRF tại OwnerAccessFilter.',
'SportFieldDB.setAutoApproval() mở transaction, tìm sân với PESSIMISTIC_WRITE, kiểm tra quyền sở hữu và loại sân được hỗ trợ; sau đó gọi setAutoApprove(enabled) và commit. Nếu lỗi thì rollback, cuối cùng đóng EntityManager.',
'Servlet lưu thông báo kết quả và redirect về trang lịch. Request GET tiếp theo nạp lại dữ liệu để JSP hiển thị chế độ mới.'
]:p('• '+t)

h('4. Tiêu chí kiểm tra và giới hạn')
table(['Tình huống','Kết quả mong đợi'],[
('Bật cho sân pickleball/bóng rổ thuộc chủ sân','Lưu autoApprove=true; trang lịch hiển thị tự xác nhận.'),
('Tắt chế độ','Lưu autoApprove=false; đơn tạo mới theo chế độ duyệt thủ công.'),
('Bật cho gym/hồ bơi hoặc sửa sân không thuộc owner','Không cho bật loại sân không hỗ trợ hoặc cập nhật sân ngoài quyền sở hữu.'),
('Đã có đơn PENDING trước khi bật','Đơn cũ vẫn PENDING; không tự động duyệt lại.')])
p('Bảng trên là tiêu chí kiểm tra, không phải biên bản xác nhận đã chạy và đạt tất cả tình huống. Báo cáo được đối chiếu với mã nguồn hiện tại; chưa kèm ảnh giao diện, log kiểm thử hoặc kết quả đọc lại database.')
p('Tự xác nhận chỉ được xét khi tạo đơn mới hợp lệ trong BookingDB.createCustomerBooking(). Đây không phải tác vụ chạy nền và không tự chuyển đơn sang COMPLETED khi hết giờ. autoApprove là cấu hình của sân, khác với PENDING/CONFIRMED là trạng thái của một đơn.')

h('5. Kết luận')
p('Chức năng có phạm vi nhỏ nhưng thể hiện luồng từ form JSP đến Servlet, kiểm tra quyền và cập nhật dữ liệu trong transaction. Nội dung cần nắm là cách nhận request parameter, lấy owner từ session, ánh xạ thuộc tính JPA và dùng redirect sau POST. Việc bật/tắt cấu hình không thay thế toàn bộ quy trình xử lý đơn đặt sân của nhóm.')

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Huỳnh Nhật Bảo — 24110079 | Trang ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Báo cáo cá nhân — Bật/tắt tự xác nhận đơn mới cho sân'
doc.core_properties.author='Huỳnh Nhật Bảo'
doc.core_properties.subject='Web Programming — Sportify'
doc.save(out)
check=Document(out)
alltext='\n'.join(x.text for x in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for row in t.rows for c in row.cells)
for s in ['Huỳnh Nhật Bảo','24110079','setAutoApproval','Đơn cũ vẫn PENDING','không bao gồm toàn bộ chức năng Booking']:
 assert s in alltext,s
with zipfile.ZipFile(out) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
report={'file':str(out),'size_bytes':out.stat().st_size,'sections':len([x for x in check.paragraphs if x.style.name=='Heading 1']),'tables':len(check.tables),'identity_verified':True,'docx_readback':'PASS','zip_xml_health':'PASS','note':'No runtime testing or visual pagination claim.'}
(outdir/'BAO_CAO_CA_NHAN_24110079_VERIFY.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
